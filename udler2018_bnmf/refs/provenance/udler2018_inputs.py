#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
udler2018_inputs.py — Udler 2018(PLoS Med) 재현용 입력 목록을 논문 supplement에서 '코드로' 구성.

사람이 대는 입력: supplement 파일 경로뿐. 그 외 판정(어떤 형질이 실제 열인지,
어느 소스가 아직 살아있는지)은 전부 코드가 계산한다.

입력 파일 (refs/):
  udler_s006.xlsx  S1 Table  변이 목록 (Position_alleles = VAR_ID 형식, hg19)
  udler_s007.xlsx  S2 Table  형질 GWAS 데이터셋 (Consortium / Website)
  udler_s008.xlsx  S3 Table  정답지: 변이 x 클러스터 가중치
  udler_s009.xlsx  S4 Table  정답지: 형질 x 클러스터 가중치  <- 실제 열 목록의 출처

출력 (outputs/):
  udler2018_variants.csv        94~95개 변이 (VAR_ID / rsID / risk allele)
  udler2018_columns.csv         S4 기준 실제 bNMF 열 (_pos/_neg 분해 전/후)
  udler2018_source_manifest.csv 형질 -> consortium -> website (website 전방 채움)
  udler2018_source_probe.csv    각 distinct URL의 현재 응답 상태

실행: python src/udler2018_inputs.py [--probe]
"""
import os, re, sys, csv, collections
import openpyxl

REFS = os.path.join(os.path.dirname(__file__), "..", "refs")
OUT  = os.path.join(os.path.dirname(__file__), "..", "outputs")


def sheet_rows(fname, sheet=None):
    wb = openpyxl.load_workbook(os.path.join(REFS, fname), read_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    return [[("" if c is None else str(c).strip()) for c in r]
            for r in ws.iter_rows(values_only=True)]


# --------------------------------------------------------------------------
# 1. S1 -> 변이 목록
# --------------------------------------------------------------------------
def parse_variants():
    """S3(정답지)의 94개 SNP을 권위 목록으로 삼고, S1에서 VAR_ID를 붙인다.

    S1 Table은 (a)분석 포함 SNP + (b)제외 SNP(사유 기재)이 한 시트에 이어 붙어 있어
    S1만 읽으면 제외분까지 섞인다. 그래서 '무엇이 행인가'는 S3가 정의하게 하고
    S1은 좌표 조회용으로만 쓴다.
    """
    s1 = sheet_rows("udler_s006.xlsx")
    hi = next(i for i, r in enumerate(s1) if "Position_alleles" in r)
    h = s1[hi]
    ip, iv = h.index("Position_alleles"), h.index("Variant")
    ira = h.index("Risk allele") if "Risk allele" in h else None
    pos = {}
    for r in s1[hi + 1:]:
        if len(r) <= ip or not r[iv].startswith("rs"):
            continue
        pos.setdefault(r[iv], (r[ip], r[ira] if ira is not None else ""))

    s3 = sheet_rows("udler_s008.xlsx")
    h3i = next(i for i, r in enumerate(s3) if "SNP" in r)
    isnp = s3[h3i].index("SNP")
    iloc = s3[h3i].index("Loci") if "Loci" in s3[h3i] else None

    out = []
    for r in s3[h3i + 1:]:
        if len(r) <= isnp or not r[isnp]:
            continue
        p, ra = pos.get(r[isnp], ("", ""))
        multi = "," in p          # 다대립 (예: rs1801282 = 3_12393125_C_G,A)
        out.append({"VAR_ID": p, "rsID": r[isnp],
                    "locus": r[iloc] if iloc is not None else "",
                    "Risk_Allele": ra,
                    "multiallelic": "YES" if multi else "",
                    "var_id_ok": "" if (multi or not p) else "YES"})
    return out


# --------------------------------------------------------------------------
# 2. S4 -> 실제 bNMF 열 목록 (정답지가 곧 열 정의)
#    S4의 행 이름은 "<trait>_pos" / "<trait>_neg" 형태.
# --------------------------------------------------------------------------
def parse_columns():
    rows = sheet_rows("udler_s009.xlsx")
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Trait")
    cols = [r[0] for r in rows[hi + 1:] if r and r[0]]
    base = collections.OrderedDict()
    for c in cols:
        m = re.match(r"^(.*)_(pos|neg)$", c)
        base.setdefault(m.group(1) if m else c, []).append(c)
    return cols, base


# --------------------------------------------------------------------------
# 3. S2 -> 소스 매니페스트 (Website는 그룹 첫 행에만 있어 전방 채움 필요)
#    Website 칸에 URL이 아닌 설명문이 들어간 행이 있어 URL 판정으로 거른다.
# --------------------------------------------------------------------------
def parse_sources():
    rows = sheet_rows("udler_s007.xlsx")
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Trait")
    h = rows[hi]
    it, iN, ic, iw = (h.index("Trait"), h.index("N"),
                      h.index("Consortium"), h.index("Website"))
    recs, last_url, section = [], "", ""
    for r in rows[hi + 1:]:
        if not any(r):
            continue
        trait = r[it]
        if not trait:
            continue
        # 섹션 헤더: Trait만 있고 N/Consortium 비어 있음
        if not r[iN] and not r[ic]:
            section = trait
            continue
        w = r[iw] if len(r) > iw else ""
        if w.startswith("http"):
            last_url = w
        elif w and not w.startswith("http"):
            # 설명문(지방산 이름 등) -> URL 아님. 직전 URL 유지.
            pass
        recs.append({"section": section, "trait": trait, "N": r[iN],
                     "consortium": r[ic], "website": last_url,
                     "website_raw": w})
    return recs


# --------------------------------------------------------------------------
# 4. 소스 URL 생존 확인 (2018년 논문이라 link rot 가능)
# --------------------------------------------------------------------------
def probe(urls):
    import requests
    out = []
    for u in urls:
        st, note = "", ""
        try:
            r = requests.get(u, timeout=30, allow_redirects=True,
                             headers={"User-Agent": "Mozilla/5.0"})
            st = str(r.status_code)
            if r.url.rstrip("/") != u.rstrip("/"):
                note = "redirected -> " + r.url
        except Exception as e:
            st, note = "ERROR", type(e).__name__ + ": " + str(e)[:120]
        out.append({"url": u, "status": st, "note": note})
        print(f"  [{st:>5}] {u}  {note}", flush=True)
    return out


def write_csv(path, rows, cols):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in cols})
    print("wrote", path)


def main():
    os.makedirs(OUT, exist_ok=True)

    v = parse_variants()
    write_csv(os.path.join(OUT, "udler2018_variants.csv"), v,
              ["VAR_ID", "rsID", "locus", "Risk_Allele", "multiallelic", "var_id_ok"])
    n_ok = sum(1 for x in v if x["var_id_ok"])
    n_multi = sum(1 for x in v if x["multiallelic"])
    n_miss = sum(1 for x in v if not x["VAR_ID"])
    print(f"변이: {len(v)}개 (VAR_ID 정상 {n_ok} / 다대립 {n_multi} / S1에 없음 {n_miss})")

    cols, base = parse_columns()
    rows = [{"bnmf_column": c} for c in cols]
    write_csv(os.path.join(OUT, "udler2018_columns.csv"), rows, ["bnmf_column"])
    print(f"bNMF 열: {len(cols)}개  /  기저 형질: {len(base)}개")

    s = parse_sources()
    write_csv(os.path.join(OUT, "udler2018_source_manifest.csv"), s,
              ["section", "trait", "N", "consortium", "website", "website_raw"])
    print(f"S2 형질 행: {len(s)}개")
    print("consortium 분포:",
          dict(collections.Counter(x["consortium"] or "(none)" for x in s)))

    if "--probe" in sys.argv:
        urls = sorted({x["website"] for x in s if x["website"]})
        print(f"\n소스 URL {len(urls)}개 확인 중...")
        p = probe(urls)
        write_csv(os.path.join(OUT, "udler2018_source_probe.csv"), p,
                  ["url", "status", "note"])
        ok = sum(1 for x in p if x["status"] == "200")
        print(f"\n생존: {ok}/{len(p)}")


if __name__ == "__main__":
    main()
