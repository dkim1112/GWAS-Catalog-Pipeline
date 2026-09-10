#!/usr/bin/env python3
"""로컬 manifest.xlsx 상태 확인."""
import openpyxl
p = "/Users/kde/Documents/GWAS Catalog Pipeline/udler2018_bnmf/manifest.xlsx"
wb = openpyxl.load_workbook(p)
ws = wb["trait_gwas"]
names = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
print(f"local trait_gwas rows: {len(names)}")
print(f"CHARGE 5 in manifest: {sorted(n for n in names if n in ('dpa','palmitoleic','n6_1821','n6_1831','n6_2031'))}")
print("all traits:", sorted(names))
