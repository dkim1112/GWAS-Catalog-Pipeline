#!/usr/bin/env python3
"""
_append_charge_manifest.py — one-off fix:
    manifest.xlsx 에 CHARGE 5개 행 추가 (sumstats_converted 파일 이미 존재 전제).

이 스크립트는 서버 udler2018_bnmf/ 안에서 실행하세요:
    cd ~/udler2018_bnmf && python3 _append_charge_manifest.py

sync_up.sh 가 manifest.xlsx 를 실수로 로컬(35)으로 덮어썼을 때 복구용.
convert 는 이미 되어 있으므로 append + validate 만 실행.
"""
import openpyxl, os

N = {
    "dpa":         8866,
    "palmitoleic": 8961,
    "n6_1821":     8631,
    "n6_1831":     8631,
    "n6_2031":     8631,
}

wb = openpyxl.load_workbook("manifest.xlsx")
ws = wb["trait_gwas"]
have = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
root = os.path.abspath(".")
added = 0

for k, n in N.items():
    conv = os.path.join(root, "sumstats_converted", f"{k}.tsv.gz")
    if k in have:
        print(f"  [skip] {k} 이미 매니페스트에 있음")
        continue
    if not os.path.exists(conv):
        print(f"  !! {k} 변환본 없음 — 건너뜀")
        continue
    ws.append([k, conv, n])
    added += 1
    print(f"  [add ] {k} (N={n})")

wb.save("manifest.xlsx")
print(f"\n{added}행 추가 -> trait_gwas 총 {ws.max_row-1}행")
