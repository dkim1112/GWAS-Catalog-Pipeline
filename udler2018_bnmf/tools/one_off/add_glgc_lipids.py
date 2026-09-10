#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
add_glgc_lipids.py — 지질 4개(hdl/ldl/tc/tg)를 GLGC Willer 2013 으로 추가.

배경
  Udler 2018 은 ENGAGE 2015 를 썼으나 배포처(diagram-consortium.org)가 죽었고
  GWAS Catalog 에도 요약통계가 없다. 같은 세대·같은 표본 자릿수인
  GLGC Willer 2013 jointGwasMc 로 대체한다. 원 논문과 출처가 다르므로
  결과 대조 시 반드시 명시할 것.

  원본 포맷 (탭 구분, hg19 좌표 + rsID 둘 다 있음):
    SNP_hg18  SNP_hg19  rsid  A1  A2  beta  se  N  P-value  Freq.A1.1000G.EUR
  convert_sumstats.py 의 컬럼 매칭에 그대로 걸리므로 변환 코드 수정은 불필요.

sample_size 에 대해
  prep_bNMF_2025.R:271-273 이 매니페스트 sample_size 로 파일의 SNP별 N 을
  **덮어쓴다**. 그리고 line 688 에서 z = z/sqrt(N) 스케일링에 쓰인다.
  따라서 추측값을 넣으면 안 되고, 변환본 N_PH 의 중앙값을 계산해 넣는다.

!! 이미 적용 완료된 일회성 스크립트입니다 (v2 인풋에 반영됨).
   다시 돌릴 일은 없고, "이 4개 열이 어떻게 들어왔는가" 의 기록으로 남겨둡니다.

사용:
  cd <프로젝트 루트> && python3 tools/one_off/add_glgc_lipids.py
  (data/sumstats/{hdl,ldl,tc,tg}.txt.gz 가 이미 내려받아져 있어야 함)
"""
import os, sys, csv, gzip, subprocess, statistics

# 이 파일은 tools/one_off/ 안에 있고, 데이터는 프로젝트 루트 기준이다.
HERE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
KEYS = ["hdl", "ldl", "tc", "tg"]
URL  = "http://csg.sph.umich.edu/willer/public/lipids2013/jointGwasMc_{}.txt.gz"
NOTE = "GLGC Willer 2013 jointGwasMc (ENGAGE 2015 대체); hg19; EUR"


def step(n, msg):
    print(f"[{n}/4] {msg}", flush=True)


# --- 1. 원본 존재 확인 -------------------------------------------------------
step(1, "원본 확인")
missing = [k for k in KEYS if not os.path.exists(f"{HERE}/data/sumstats/{k}.txt.gz")]
if missing:
    sys.exit(f"  !! data/sumstats/ 에 없음: {missing}\n"
             f"     먼저 내려받으세요: curl -o data/sumstats/<key>.txt.gz "
             f"{URL.format('<KEY>')}")
for k in KEYS:
    mb = os.path.getsize(f"{HERE}/data/sumstats/{k}.txt.gz") / 1e6
    print(f"  {k:4s} {mb:7.1f} MB")


# --- 2. inputs_manifest.csv 갱신 --------------------------------------------
step(2, "inputs_manifest.csv 갱신")
man_path = f"{HERE}/inputs_manifest.csv"
with open(man_path) as f:
    rd = csv.DictReader(f)
    fields, rows = rd.fieldnames, list(rd)

n_upd = 0
for r in rows:
    if r["key"] in KEYS:
        r["source"]     = "GLGC Willer 2013"
        r["status"]     = "DIRECT_URL"
        r["direct_url"] = URL.format(r["key"].upper())
        r["portal_url"] = "http://csg.sph.umich.edu/willer/public/lipids2013/"
        r["note"]       = NOTE
        n_upd += 1
with open(man_path, "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=fields)
    w.writeheader(); w.writerows(rows)
print(f"  {n_upd}개 행 갱신")


# --- 3. 포맷 변환 ------------------------------------------------------------
step(3, "포맷 변환 (4개, 각 10분 내외)")
rc = subprocess.call([sys.executable, f"{HERE}/tools/convert_sumstats.py"] + KEYS, cwd=HERE)
if rc != 0:
    sys.exit(f"  !! convert_sumstats.py 실패 (rc={rc})")


# --- 4. manifest.xlsx 에 추가 (N 은 변환본에서 실측) --------------------------
step(4, "manifest.xlsx 추가 + 점검")
import openpyxl

def median_n(path, cap=500_000):
    """변환본 N_PH 의 중앙값. prep_bNMF 의 sample size 스케일링에 쓰이는 값."""
    vals = []
    with gzip.open(path, "rt") as f:
        rd = csv.DictReader(f, delimiter="\t")
        for i, row in enumerate(rd):
            if i >= cap:
                break
            try:
                vals.append(float(row["N_PH"]))
            except (TypeError, ValueError):
                pass
    if not vals:
        return None
    return int(round(statistics.median(vals)))

wb = openpyxl.load_workbook(f"{HERE}/manifest.xlsx")
ws = wb["trait_gwas"]
have = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
added = 0

for k in KEYS:
    conv = f"{HERE}/data/sumstats_converted/{k}.tsv.gz"
    if k in have:
        print(f"  [skip] {k} 이미 매니페스트에 있음")
        continue
    if not os.path.exists(conv):
        print(f"  !! {k} 변환본 없음 — 건너뜀")
        continue
    n = median_n(conv)
    if n is None:
        print(f"  !! {k} N_PH 를 읽지 못함 — 건너뜀 (수동 확인 필요)")
        continue
    ws.append([k, conv, n])
    added += 1
    print(f"  [add ] {k}  N(중앙값)={n:,}")

wb.save(f"{HERE}/manifest.xlsx")
print(f"\n  {added}행 추가 -> trait_gwas 총 {ws.max_row - 1}행")

print()
subprocess.call([sys.executable, f"{HERE}/tools/validate_inputs.py"], cwd=HERE)

print("""
------------------------------------------------------------------
다음 단계
형질이 늘었으므로 z 행렬이 바뀝니다. 파이프라인을 다시 돌리세요.
변이 목록(pruned_vars)은 udler_substitute.R 이 매번
refs/udler2018_variants.csv 에서 새로 만들므로 재실행 부담 없음.

  Rscript run_pipeline.R

v1(35열) 결과는 results/udler2018_eur_v1/ 에 있습니다 (대조용).
------------------------------------------------------------------""")
